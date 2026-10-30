"""Check editable artifacts, traceability, PDF text and absence of configured secrets."""

import hashlib
import json
import subprocess
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from report_content import CASES, HU, RF, RNF

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "informe y rúbrica/corregido"
EV = ROOT / ".scratch/cumplimiento-rubrica-final/evidencias"


def main():
    checks = {}
    docpath = OUT / "APF1 - Pruebas de Software - Grupo 1 - corregido.docx"
    doc = Document(docpath)
    doc_text = "\n".join(
        [p.text for p in doc.paragraphs]
        + [c.text for t in doc.tables for row in t.rows for c in row.cells]
    )
    wb = load_workbook(OUT / "Tablas del informe - corregidas.xlsx")
    pres = Presentation(OUT / "Sustentacion - Grupo 1.pptx")
    checks["editable_files_open"] = True
    checks["thirteen_workbook_sheets"] = len(wb.sheetnames) == 13
    from test_inventory import TYPES
    from report_content import CASE_TEST_TYPES
    plan = wb["06 Plan de pruebas"]
    checks["separate_types_sheet_removed"] = "14 Tipos de pruebas" not in wb.sheetnames
    checks["type_column_named"] = plan["D1"].value == "Tipo de prueba"
    checks["all_36_plans_classified"] = plan.max_row == 37 and all(
        plan.cell(i, 4).value == CASE_TEST_TYPES[plan.cell(i, 1).value]
        and plan.cell(i, 4).value in TYPES for i in range(2, 38)
    )
    checks["thirty_six_complete_cases"] = wb[
        "06 Plan de pruebas"
    ].max_row == 37 and all(
        all(c.value is not None for c in row)
        for row in wb["06 Plan de pruebas"].iter_rows(min_row=2)
    )
    checks["ten_complete_stories"] = wb["07 Historias de usuario"].max_row == 11
    checks["all_ids_in_word"] = all(row[0] in doc_text for row in CASES + HU + RF + RNF)
    case_ids = {row[0] for row in CASES}
    hu_ids = {row[0] for row in HU}
    checks["unique_ids"] = len(case_ids) == len(CASES) and len(hu_ids) == len(HU)
    checks["all_requirements_have_existing_cases"] = all(
        set(r[5].split(", ")) <= case_ids for r in RF + RNF
    )
    checks["all_cases_have_existing_stories"] = all(
        set(c[1].split(", ")) <= hu_ids for c in CASES
    )
    checks["twelve_slides_with_notes"] = len(pres.slides) == 12 and all(
        sl.notes_slide.notes_text_frame.text.strip() for sl in pres.slides
    )
    pdf = OUT / "APF1 - Pruebas de Software - Grupo 1 - corregido.pdf"
    pdf_text = subprocess.check_output(["pdftotext", str(pdf), "-"], text=True)
    checks["all_cases_in_pdf"] = all(c[0] in pdf_text for c in CASES)
    checks["pdf_has_references_and_conclusions"] = (
        "Conclusiones" in pdf_text and "Referencias" in pdf_text
    )
    results = json.loads((EV / "resultados-casos.json").read_text())
    execution = {wb["08 Ejecución"].cell(i, 1).value: wb["08 Ejecución"].cell(i, 4).value for i in range(2, 38)}
    checks["execution_statuses_match_evidence"] = all(execution.get(r["id"]) == r["status"] for r in results)
    evidence_missing = []
    for r in results:
        for item in r["evidence"].split(";"):
            name = item.strip().split("#")[0]
            if name and name != "artefactos.json" and not (EV / name).exists():
                evidence_missing.append(name)
    checks["evidence_paths_resolve"] = not evidence_missing
    originals = ROOT / "informe y rúbrica"
    hashpath = EV / "originales-sha256.json"
    hashes = {
        p.name: hashlib.sha256(p.read_bytes()).hexdigest()
        for p in originals.iterdir()
        if p.is_file()
    }
    if hashpath.exists():
        checks["originals_unchanged_since_record"] = (
            json.loads(hashpath.read_text()) == hashes
        )
    else:
        hashpath.write_text(json.dumps(hashes, ensure_ascii=False, indent=2) + "\n")
        checks["originals_recorded"] = True
    secrets = []
    for name in [".moodle.env", "rag-service/.env"]:
        for line in (ROOT / name).read_text().splitlines():
            if "=" not in line or line.lstrip().startswith("#"):
                continue
            key, value = line.split("=", 1)
            value = value.strip().strip("\"'")
            if (
                any(s in key.upper() for s in ["KEY", "SECRET", "PASSWORD"])
                and len(value) >= 12
            ):
                secrets.append(value.encode())
    leaked = []
    for path in list(OUT.rglob("*")) + list(EV.rglob("*")):
        if not path.is_file() or path.suffix.lower() in [".png", ".pdf", ".svg"]:
            continue
        if path.suffix in [".docx", ".xlsx", ".pptx"]:
            with ZipFile(path) as z:
                content = b"\n".join(
                    z.read(n) for n in z.namelist() if n.endswith(".xml")
                )
        else:
            content = path.read_bytes()
        if any(secret in content for secret in secrets):
            leaked.append(str(path.relative_to(ROOT)))
    checks["configured_secrets_absent"] = not leaked
    result = {
        "passed": all(checks.values()),
        "timestamp_utc": datetime.now(timezone.utc).isoformat(),
        "summary": "Archivos editables abren; IDs y referencias verificadas; PDF conserva casos; PPT contiene notas; secretos configurados ausentes. La inspección visual se registra por separado.",
        "checks": checks,
        "missing_evidence": sorted(set(evidence_missing)),
        "secret_leak_paths": leaked,
        "pdf_info": subprocess.check_output(["pdfinfo", str(pdf)], text=True),
    }
    (EV / "artefactos.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2) + "\n"
    )
    print(
        json.dumps({"passed": result["passed"], "checks": checks}, ensure_ascii=False)
    )
    return 0 if result["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
